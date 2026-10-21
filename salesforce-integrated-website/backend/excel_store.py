from pathlib import Path
from datetime import datetime
from openpyxl import Workbook, load_workbook

DATA_DIR = Path(__file__).resolve().parent / "data"
DATABASE_FILE = DATA_DIR / "college_data.xlsx"

SHEETS = {
    "Users": ["id", "name", "email", "passwordHash", "role"],
    "FAQs": ["id", "question", "answer", "category"],
    "Queries": ["id", "name", "email", "message", "status", "created_at", "phone"],
    "Announcements": ["id", "title", "message", "category", "created_at", "status"],
}

DEFAULT_FAQS = [
    [1, "What is Salesforce?", "Salesforce is a cloud-based customer relationship management (CRM) platform used to manage customer information, sales, service, and business workflows.", "General"],
    [2, "What is CRM?", "CRM stands for Customer Relationship Management. It helps an organization manage interactions with customers and potential customers.", "General"],
    [3, "How do I register for an account?", "Open the registration form, enter your name, email address, and password, then submit the form.", "Account"],
    [4, "How do I log in?", "Open the login form and enter the email address and password you used when registering.", "Account"],
    [5, "I forgot my password. What should I do?", "Use the password recovery option if your organization has enabled it. Otherwise, contact your system administrator.", "Account"],
    [6, "How do I log out?", "Use the Logout option in the website navigation or account menu.", "Account"],
    [7, "What is an account in Salesforce?", "An account represents a company or organization you do business with. The exact fields depend on your Salesforce configuration.", "Accounts"],
    [8, "How do I create an account record?", "Open Accounts, choose New, complete the required fields, and select Save. This demo dashboard may need Salesforce API integration to store real records.", "Accounts"],
    [9, "How do I update account information?", "Open the account record, choose Edit, update the required fields, and save your changes.", "Accounts"],
    [10, "What is a contact?", "A contact represents an individual associated with a business account.", "Accounts"],
    [11, "What is a lead?", "A lead is a person or organization that may become a potential customer.", "Sales"],
    [12, "What is an opportunity?", "An opportunity represents a potential sales deal, including its value, stage, and expected close date.", "Sales"],
    [13, "How do I create a task?", "Open Tasks, choose New, enter the subject, due date, and other details, then save.", "Tasks"],
    [14, "How do I mark a task as complete?", "Open the task and change its status to Completed, or use the completion control provided by your configured interface.", "Tasks"],
    [15, "What are Salesforce reports?", "Reports organize, filter, and summarize records so users can analyze business data.", "Reports"],
    [16, "How do I open a report?", "Choose Reports in the navigation and select the report you want to view.", "Reports"],
    [17, "What is a dashboard?", "A dashboard displays visual summaries, such as charts and key metrics, often based on reports.", "Dashboards"],
    [18, "Why is my dashboard empty?", "A dashboard may be empty if no records match its filters, the source report has no data, or permissions prevent access.", "Dashboards"],
    [19, "How do I search FAQs?", "Open FAQ & Help and enter a word or phrase in the search field. You can also filter by category.", "FAQ"],
    [20, "How can I contact support?", "Submit a support query through the website if that feature is enabled, or contact your organization's help desk.", "Support"],
    [21, "Why are my FAQs not loading?", "Check that the FastAPI backend is running and that the API endpoint /api/faqs responds successfully.", "Technical Support"],
    [22, "Where are FAQ questions stored?", "In this project, FAQ questions and answers are stored in the FAQs sheet of backend/data/college_data.xlsx.", "Database"],
    [23, "How do I add a new FAQ?", "Open college_data.xlsx, add a row to the FAQs sheet using the existing columns, save the workbook, and refresh the FAQ page. Use a unique ID.", "Database"],
    [24, "Can I add FAQ categories?", "Yes. Enter the category name in the category column of the FAQs sheet. Your frontend must support the category for it to appear as a filter option.", "Database"],
    [25, "What does the Queries sheet store?", "The Queries sheet stores submitted support messages and their status when query submission is implemented.", "Database"],
    [26, "What does the Announcements sheet store?", "The Announcements sheet stores notices or updates that can be displayed on the website when the relevant frontend feature is implemented.", "Database"],
    [27, "Is this connected to a live Salesforce org?", "Not automatically. This project uses a Salesforce-inspired interface and an Excel-backed FastAPI demo. Live Salesforce records require Salesforce API authentication and integration.", "Integration"],
    [28, "How do I report a technical problem?", "Describe the problem, include any error message, and submit it through the support process available to your organization.", "Support"],
    [29, "Why do I see a CORS error?", "The backend must allow the exact origin used by your frontend. Check FRONTEND_URL in the backend environment configuration and restart FastAPI.", "Technical Support"],
    [30, "How do I start the backend?", "Open a terminal in the backend folder, activate the intended Python environment, install requirements, then run: uvicorn main:app --reload.", "Technical Support"],
]

def ensure_workbook():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if not DATABASE_FILE.exists():
        wb = Workbook()
        default = wb.active
        wb.remove(default)
        for sheet_name, headers in SHEETS.items():
            sheet = wb.create_sheet(sheet_name)
            sheet.append(headers)
            if sheet_name == "FAQs":
                for row in DEFAULT_FAQS:
                    sheet.append(row)
        wb.save(DATABASE_FILE)
        wb.close()
    else:
        # Add any missing sheets without overwriting existing data.
        wb = load_workbook(DATABASE_FILE)
        changed = False
        for sheet_name, headers in SHEETS.items():
            if sheet_name not in wb.sheetnames:
                sheet = wb.create_sheet(sheet_name)
                sheet.append(headers)
                if sheet_name == "FAQs":
                    for row in DEFAULT_FAQS:
                        sheet.append(row)
                changed = True
        if "FAQs" in wb.sheetnames and wb["FAQs"].max_row <= 1:
            for row in DEFAULT_FAQS:
                wb["FAQs"].append(row)
            changed = True
        if changed:
            wb.save(DATABASE_FILE)
        wb.close()

def read_sheet(sheet_name: str) -> list[dict]:
    ensure_workbook()
    wb = load_workbook(DATABASE_FILE, read_only=True, data_only=True)
    try:
        sheet = wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    result = []
    for row in rows[1:]:
        if not any(v is not None for v in row):
            continue
        item = {headers[i]: (row[i] if i < len(row) and row[i] is not None else "")
                for i in range(len(headers)) if headers[i]}
        result.append(item)
    return result

def append_row(sheet_name: str, values: dict):
    ensure_workbook()
    wb = load_workbook(DATABASE_FILE)
    try:
        sheet = wb[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        sheet.append([values.get(header, "") for header in headers])
        wb.save(DATABASE_FILE)
    finally:
        wb.close()

def read_users() -> list[dict]:
    return read_sheet("Users")

def write_users(users: list[dict]):
    ensure_workbook()
    wb = load_workbook(DATABASE_FILE)
    try:
        sheet = wb["Users"]
        sheet.delete_rows(2, max(sheet.max_row - 1, 1))
        for item in users:
            sheet.append([item.get(header, "") for header in SHEETS["Users"]])
        wb.save(DATABASE_FILE)
    finally:
        wb.close()

def read_faqs() -> list[dict]:
    return read_sheet("FAQs")

def read_queries() -> list[dict]:
    return read_sheet("Queries")

def add_query(name: str, email: str, message: str, phone: str = "") -> dict:
    """Append a support request to the Queries sheet in the Excel workbook."""
    queries = read_queries()
    next_id = max((int(item.get("id") or 0) for item in queries), default=0) + 1
    item = {
        "id": next_id,
        "name": name.strip(),
        "email": email.strip().lower(),
        "message": message.strip(),
        "status": "Open",
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "phone": phone.strip(),
    }
    append_row("Queries", item)
    return item

def read_announcements() -> list[dict]:
    return read_sheet("Announcements")

def add_faq(question: str, answer: str, category: str) -> dict:
    faqs = read_faqs()
    next_id = max((int(item.get("id") or 0) for item in faqs), default=0) + 1
    item = {"id": next_id, "question": question.strip(), "answer": answer.strip(), "category": category.strip() or "General"}
    append_row("FAQs", item)
    return item

# Backward-compatible name used by main.py.
def ensure_workbooks():
    ensure_workbook()


def update_query_status(query_id: int, status: str) -> dict | None:
    """Update a support query status in the Queries sheet."""
    allowed_statuses = {"Open", "In Progress", "Resolved"}
    if status not in allowed_statuses:
        raise ValueError("Status must be Open, In Progress, or Resolved.")
    ensure_workbook()
    wb = load_workbook(DATABASE_FILE)
    try:
        sheet = wb["Queries"]
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
        if "id" not in headers or "status" not in headers:
            raise ValueError("Queries sheet must contain id and status columns.")
        id_col = headers.index("id") + 1
        status_col = headers.index("status") + 1
        for row_num in range(2, sheet.max_row + 1):
            try:
                current_id = int(sheet.cell(row=row_num, column=id_col).value)
            except (TypeError, ValueError):
                continue
            if current_id == query_id:
                sheet.cell(row=row_num, column=status_col).value = status
                wb.save(DATABASE_FILE)
                return {
                    headers[col - 1]: (sheet.cell(row=row_num, column=col).value or "")
                    for col in range(1, len(headers) + 1)
                }
        return None
    finally:
        wb.close()
